"""Document and source-code ingestion.

Parses engineering inputs into page-anchored blocks, then groups them into
section-aware chunks that keep the heading path and page range needed for
citations.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from app.config import settings

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm", ".csv"}
CODE_EXTENSIONS = {
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".cc": "cpp",
    ".hpp": "cpp",
    ".py": "python",
    ".java": "java",
    ".cs": "csharp",
    ".capl": "capl",
    ".can": "capl",
    ".m": "matlab",
    ".arxml": "xml",
    ".xml": "xml",
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
}
SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS | set(CODE_EXTENSIONS)

# A heading is a numbered clause ("4.2.1 Signal Routing"), an ALL CAPS line, or
# a markdown heading. Engineering PDFs overwhelmingly use the numbered form.
_NUMBERED_HEADING = re.compile(r"^\s{0,6}(\d{1,2}(?:\.\d{1,2}){0,4})[.)]?\s+(\S.{0,110})$")
_MARKDOWN_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(\S.*)$")
_UPPER_HEADING = re.compile(r"^\s{0,6}([A-Z][A-Z0-9 ,/&()\-]{5,70})\s*$")


@dataclass
class Block:
    """A unit of text tied to a physical location in the source."""

    text: str
    page: int
    kind: str = "text"  # text | table | ocr | code


@dataclass
class Chunk:
    """An embeddable chunk carrying everything needed to cite it."""

    chunk_id: str
    text: str
    ordinal: int
    page_start: int
    page_end: int
    heading_path: str
    kind: str
    extra: dict = field(default_factory=dict)

    def citation_label(self) -> str:
        pages = (
            f"p.{self.page_start}"
            if self.page_start == self.page_end
            else f"pp.{self.page_start}-{self.page_end}"
        )
        return f"{self.heading_path or 'Document'} ({pages})" if self.heading_path else pages

    def to_metadata(self) -> dict:
        meta = {
            "ordinal": self.ordinal,
            "page_start": self.page_start,
            "page_end": self.page_end,
            "heading_path": self.heading_path,
            "kind": self.kind,
        }
        for key, value in self.extra.items():
            if isinstance(value, (str, int, float, bool)):
                meta[key] = value
        return meta


class IngestionError(RuntimeError):
    pass


# --------------------------------------------------------------------------
# Parsers
# --------------------------------------------------------------------------


def _ocr_page(page) -> str:
    """Render a PDF page to an image and OCR it. Returns '' if unavailable."""
    if not settings.enable_ocr:
        return ""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    if settings.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
    try:
        import io

        pixmap = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        return pytesseract.image_to_string(image).strip()
    except Exception:
        # OCR is best-effort; a missing Tesseract binary must not fail ingestion.
        return ""


def _table_to_text(table: Sequence[Sequence[str | None]]) -> str:
    rows = []
    for row in table:
        cells = [(cell or "").replace("\n", " ").strip() for cell in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_pdf(path: Path) -> list[Block]:
    import fitz  # PyMuPDF

    blocks: list[Block] = []
    with fitz.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < 40:
                ocr_text = _ocr_page(page)
                if ocr_text:
                    blocks.append(Block(text=ocr_text, page=index, kind="ocr"))
                    continue
            if text:
                blocks.append(Block(text=text, page=index, kind="text"))

    # Tables are extracted separately so their structure survives chunking.
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                for table in page.extract_tables() or []:
                    rendered = _table_to_text(table)
                    if len(rendered) > 30:
                        blocks.append(
                            Block(text=f"[TABLE]\n{rendered}", page=index, kind="table")
                        )
    except Exception:
        pass

    if not blocks:
        raise IngestionError(f"No extractable content found in {path.name}")
    blocks.sort(key=lambda b: (b.page, 0 if b.kind != "table" else 1))
    return blocks


def parse_docx(path: Path) -> list[Block]:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            level = "".join(ch for ch in paragraph.style.name if ch.isdigit()) or "1"
            parts.append(f"{'#' * int(level)} {text}")
        else:
            parts.append(text)
    for table in document.tables:
        rendered = _table_to_text([[cell.text for cell in row.cells] for row in table.rows])
        if rendered:
            parts.append(f"[TABLE]\n{rendered}")
    if not parts:
        raise IngestionError(f"No extractable content found in {path.name}")
    # DOCX has no stable page model; treat the whole file as page 1.
    return [Block(text="\n".join(parts), page=1, kind="text")]


def parse_spreadsheet(path: Path) -> list[Block]:
    import pandas as pd

    blocks: list[Block] = []
    if path.suffix.lower() == ".csv":
        frames = {"Sheet1": pd.read_csv(path, dtype=str).fillna("")}
    else:
        frames = pd.read_excel(path, sheet_name=None, dtype=str)
    for sheet_index, (name, frame) in enumerate(frames.items(), start=1):
        frame = frame.fillna("")
        rendered = frame.to_csv(sep="|", index=False)
        blocks.append(Block(text=f"# {name}\n[TABLE]\n{rendered}", page=sheet_index, kind="table"))
    if not blocks:
        raise IngestionError(f"No sheets found in {path.name}")
    return blocks


def parse_plaintext(path: Path) -> list[Block]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        raise IngestionError(f"{path.name} is empty")
    return [Block(text=text, page=1, kind="text")]


def parse_source_file(path: Path) -> list[Block]:
    """Source code is kept line-anchored: 'page' carries the starting line."""
    text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        raise IngestionError(f"{path.name} is empty")
    return [Block(text=text, page=1, kind="code")]


def parse_file(path: Path) -> list[Block]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in {".xlsx", ".xlsm", ".csv"}:
        return parse_spreadsheet(path)
    if suffix in {".txt", ".md"}:
        return parse_plaintext(path)
    if suffix in CODE_EXTENSIONS:
        return parse_source_file(path)
    raise IngestionError(f"Unsupported file type: {suffix}")


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------


def detect_heading(line: str) -> tuple[int, str] | None:
    """Return (level, title) when the line looks like a section heading."""
    stripped = line.strip()
    if not stripped or len(stripped) > 130:
        return None

    md = _MARKDOWN_HEADING.match(line)
    if md:
        return len(md.group(1)), md.group(2).strip()

    numbered = _NUMBERED_HEADING.match(line)
    if numbered:
        number, title = numbered.groups()
        # Reject sentences that merely start with a number.
        if title.endswith((".", ";", ",")) and len(title.split()) > 14:
            return None
        return number.count(".") + 1, f"{number} {title.strip()}"

    upper = _UPPER_HEADING.match(line)
    if upper and len(stripped.split()) <= 10:
        return 1, upper.group(1).strip()
    return None


def _heading_path(stack: list[tuple[int, str]]) -> str:
    return " > ".join(title for _, title in stack)


def _make_chunk_id(document_id: str, ordinal: int) -> str:
    digest = hashlib.sha1(f"{document_id}:{ordinal}".encode()).hexdigest()[:12]
    return f"{document_id}-{ordinal:04d}-{digest}"


def chunk_blocks(
    blocks: Iterable[Block],
    document_id: str,
    target_chars: int | None = None,
    overlap_chars: int | None = None,
) -> list[Chunk]:
    """Group blocks into chunks, splitting on headings and size limits."""
    target = target_chars or settings.chunk_target_chars
    overlap = overlap_chars or settings.chunk_overlap_chars

    chunks: list[Chunk] = []
    stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    buffer_len = 0
    page_start = 1
    page_end = 1
    current_path = ""
    ordinal = 0

    def flush() -> None:
        nonlocal buffer, buffer_len, ordinal, page_start
        text = "\n".join(buffer).strip()
        if len(text) < 25:
            buffer = []
            buffer_len = 0
            return
        chunks.append(
            Chunk(
                chunk_id=_make_chunk_id(document_id, ordinal),
                text=text,
                ordinal=ordinal,
                page_start=page_start,
                page_end=page_end,
                heading_path=current_path,
                kind="text",
            )
        )
        ordinal += 1
        if overlap > 0 and len(text) > overlap:
            tail = text[-overlap:]
            buffer = [tail]
            buffer_len = len(tail)
        else:
            buffer = []
            buffer_len = 0
        page_start = page_end

    for block in blocks:
        # Tables and OCR output stay whole so their rows are never half-cited.
        if block.kind in {"table", "ocr"}:
            flush()
            chunks.append(
                Chunk(
                    chunk_id=_make_chunk_id(document_id, ordinal),
                    text=block.text.strip(),
                    ordinal=ordinal,
                    page_start=block.page,
                    page_end=block.page,
                    heading_path=current_path,
                    kind=block.kind,
                )
            )
            ordinal += 1
            page_start = page_end = block.page
            continue

        page_end = block.page
        if not buffer:
            page_start = block.page

        for line in block.text.splitlines():
            heading = detect_heading(line)
            if heading:
                flush()
                level, title = heading
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, title))
                current_path = _heading_path(stack)
                buffer = [line.strip()]
                buffer_len = len(line)
                page_start = block.page
                continue

            buffer.append(line)
            buffer_len += len(line) + 1
            if buffer_len >= target:
                flush()

    flush()
    return chunks


def chunk_code(text: str, document_id: str, language: str) -> list[Chunk]:
    """Split source code on top-level function boundaries, keeping line numbers."""
    lines = text.splitlines()
    boundaries: list[int] = [0]
    if language in {"c", "cpp"}:
        pattern = re.compile(r"^[A-Za-z_][\w \t\*\(\),]*\)\s*\{\s*$")
    elif language == "python":
        pattern = re.compile(r"^(def |class |async def )")
    elif language == "java" or language == "csharp":
        pattern = re.compile(r"^\s{0,8}(public|private|protected|static).*\)\s*\{\s*$")
    else:
        pattern = None

    if pattern is not None:
        for index, line in enumerate(lines):
            if pattern.match(line) and index not in boundaries:
                boundaries.append(index)
    boundaries.append(len(lines))
    boundaries = sorted(set(boundaries))

    chunks: list[Chunk] = []
    ordinal = 0
    for start, end in zip(boundaries, boundaries[1:]):
        segment = "\n".join(lines[start:end]).strip()
        if len(segment) < 20:
            continue
        # Oversized segments are split again so embeddings stay meaningful.
        for offset in range(0, max(len(segment), 1), 4000):
            piece = segment[offset : offset + 4000]
            if len(piece) < 20:
                continue
            chunks.append(
                Chunk(
                    chunk_id=_make_chunk_id(document_id, ordinal),
                    text=piece,
                    ordinal=ordinal,
                    page_start=start + 1,
                    page_end=end,
                    heading_path=f"lines {start + 1}-{end}",
                    kind="code",
                    extra={"language": language},
                )
            )
            ordinal += 1
    return chunks


def ingest_path(path: Path, document_id: str) -> list[Chunk]:
    """Parse and chunk a single file."""
    suffix = path.suffix.lower()
    if suffix in CODE_EXTENSIONS and suffix not in {".xml", ".json", ".yaml", ".yml"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return chunk_code(text, document_id, CODE_EXTENSIONS[suffix])
    blocks = parse_file(path)
    return chunk_blocks(blocks, document_id)


def file_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(65536), b""):
            digest.update(block)
    return digest.hexdigest()[:32]
